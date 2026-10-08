# -*- coding: utf-8 -*-
"""
knowledge_base.extractors —— 把课件/文档变成“可检索的知识点块”

处理链路：
  原始文件/PDF/PPT/Word/网页正文
    → 按页/幻灯片/段落切块（保留“第3页/第12张”定位信息）
    → 规则识别：定义、公式、例题、概念
    → 关键词（词频 + 中文分词，jieba 可选）
    → 可选 AI 知识点提炼（ai.prompts.SYSTEM_KB）
"""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Sequence

from utils.logger import get_logger
from utils.text import clean_text, normalize_for_match, one_line, tokenize

log = get_logger("knowledge_base.extractors")

FORMULA_HINT_RE = re.compile(
    r"(=|＝|≈|≤|≥|∝|√|∫|∑|⇒|Δ|Φ|φ|Ω|ω|π|μ|×|÷|\frac|\\sum|U\s*=|I\s*=|R\s*=|P\s*=|E\s*=)"
)
DEF_HINT_RE = re.compile(r"(定义|概念|定理|定律|法则|性质|条件|充要|称为|是指|即|规定)")
EXAMPLE_HINT_RE = re.compile(r"(例\s*\d|例题|示例|练习|习题|解：|答：)")
KEY_SENTENCE_RE = re.compile(r"(戴维南|诺顿|叠加|基尔霍夫|KCL|KVL|欧姆|功率因数|三相|互感|谐振|RC|RL|RLC"
                             r"|等效|阻抗|导纳|节点|网孔|回路|电容|电感|磁路|变压器|二极管|三极管|运放)")
STOPWORDS = set(
    """的 了 和 与 及 或 是 在 有 为 也 都 而 被 把 其 这 那 之 并 等 各 由 从 到 中 上 下 个 项 图 表 页 第 一 二 三 四 五
    the a an of and or to in is are was were for with on this that it as at be by not but can will would could
    我们 你们 它们 可以 进行 通过 使用 需要 能够 以下 如图 所示 如果 那么 因此 所以 但是 然而 例如 什么 怎么""".split()
)
MIN_DOC_CHARS = 20


# 学习通页面里的导航/目录/提示文字不是课程资料，不能进知识库
_NOISE_WORDS = (
    "请输入验证码", "手机扫码", "待完成任务点", "任务点未完成", "已完成任务点",
    "退出登录", "进入空间", "体验新版", "点击下载", "功能升级通知", "课程门户",
    "当前章节还有", "是否去完成", "知道了", "标记智慧课程", "解除智慧课程",
    "此内容由AI生成", "完成条件", "观看时长需", "不可拖拽", "页面最小化",
    "中文 (简体)", "中文 (繁体)", "任务点已完成", "重播", "倍速", "清晰度",
)
_INDEX_HEAD = re.compile(r"^\s*\d+\.\d+\s+\S")
_INDEX_ANY = re.compile(r"\d+\.\d+\s+[\u4e00-\u9fff]")
_PURE_NAV = re.compile(r"^(目录|状态|更多|块\d+|章节测验\s*\d*)$")


def is_noise(text: str) -> bool:
    """判断一段文本是否属于导航/目录噪声。"""
    s = one_line(text)
    if not s:
        return True
    if any(w in s for w in _NOISE_WORDS):
        return True
    if _PURE_NAV.match(s):
        return True
    if len(_INDEX_ANY.findall(s)) >= 3:        # 目录页：连续多条“x.y 标题”
        return True
    if _INDEX_HEAD.match(s) and len(s) < 40:   # 单独的目录条目
        return True
    return False

@dataclass
class Block:
    text: str
    loc: str = ""       # 第3页 / 第12张幻灯片 / 段落12
    kind: str = "text"  # text|heading|formula|definition|example


@dataclass
class ParsedDoc:
    title: str
    doc_type: str
    blocks: list[Block] = field(default_factory=list)
    error: str = ""

    @property
    def chars(self) -> int:
        return sum(len(b.text) for b in self.blocks)


# ------------------------------------------------------------------ 文件解析
def parse_file(path: str | Path) -> ParsedDoc:
    p = Path(path)
    suffix = p.suffix.lower()
    try:
        if suffix == ".pdf":
            return _parse_pdf(p)
        if suffix in (".docx", ".doc"):
            return _parse_docx(p)
        if suffix in (".pptx", ".ppt"):
            return _parse_pptx(p)
        if suffix in (".txt", ".md", ".markdown"):
            return _parse_text(p)
        if suffix in (".html", ".htm"):
            return _parse_html(p)
    except ImportError as exc:
        return ParsedDoc(title=p.name, doc_type=suffix.lstrip("."), error=f"缺少解析库：{exc}")
    except Exception as exc:
        log.exception("解析失败：%s", p)
        return ParsedDoc(title=p.name, doc_type=suffix.lstrip("."), error=str(exc))
    return ParsedDoc(title=p.name, doc_type=suffix.lstrip("."), error="暂不支持的文件类型")


def _parse_pdf(p: Path) -> ParsedDoc:
    from pypdf import PdfReader

    reader = PdfReader(str(p))
    blocks: list[Block] = []
    for i, page in enumerate(reader.pages, 1):
        try:
            text = clean_text(page.extract_text() or "")
        except Exception as exc:
            log.debug("PDF 第 %s 页解析失败：%s", i, exc)
            continue
        if text:
            blocks.append(Block(text=text, loc=f"第{i}页"))
    return ParsedDoc(title=p.stem, doc_type="pdf", blocks=blocks)


def _parse_docx(p: Path) -> ParsedDoc:
    from docx import Document

    doc = Document(str(p))
    blocks = [Block(text=clean_text(par.text), loc=f"段落{i}", kind="heading"
                     if (par.style and "Head" in str(par.style.name)) else "text")
              for i, par in enumerate(doc.paragraphs, 1) if one_line(par.text)]
    for t_idx, table in enumerate(doc.tables, 1):
        for r_idx, row in enumerate(table.rows):
            cells = [clean_text(c.text) for c in row.cells if one_line(c.text)]
            if cells:
                blocks.append(Block(text=" | ".join(cells), loc=f"表{t_idx}第{r_idx}行"))
    return ParsedDoc(title=p.stem, doc_type="docx", blocks=blocks)


def _parse_pptx(p: Path) -> ParsedDoc:
    from pptx import Presentation

    prs = Presentation(str(p))
    blocks: list[Block] = []
    for i, slide in enumerate(prs.slides, 1):
        texts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for para in shape.text_frame.paragraphs:
                    t = clean_text("".join(run.text for run in para.runs))
                    if t:
                        texts.append(t)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [clean_text(c.text) for c in row.cells if one_line(c.text)]
                    if cells:
                        texts.append(" | ".join(cells))
        try:
            notes = slide.notes_slide.notes_text_frame.text
        except Exception:
            notes = ""
        if one_line(notes):
            texts.append("【备注】" + clean_text(notes))
        if texts:
            blocks.append(Block(text="\n".join(texts), loc=f"第{i}张幻灯片"))
    return ParsedDoc(title=p.stem, doc_type="pptx", blocks=blocks)


def _parse_text(p: Path) -> ParsedDoc:
    raw = p.read_text(encoding="utf-8", errors="ignore")
    blocks = [Block(text=clean_text(t), loc=f"段{i}") for i, t in enumerate(re.split(r"\n\s*\n", raw), 1)
              if one_line(t)]
    if not blocks and one_line(raw):
        blocks = [Block(text=clean_text(raw), loc="全文")]
    return ParsedDoc(title=p.stem, doc_type="md" if p.suffix.lower() in (".md", ".markdown") else "txt",
                     blocks=blocks)


def _parse_html(p: Path) -> ParsedDoc:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(p.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    blocks = [Block(text=clean_text(el.get_text()), loc=f"{el.name}{i}")
              for i, el in enumerate(soup.find_all(["h1", "h2", "h3", "p", "li"]), 1) if one_line(el.get_text())]
    return ParsedDoc(title=p.stem, doc_type="html", blocks=blocks)


def parse_web_blocks(raw_blocks: Sequence[dict[str, Any]], title: str = "网页") -> ParsedDoc:
    blocks = [Block(text=clean_text(b.get("text", "")), loc=f"块{i}",
                    kind="heading" if b.get("kind") == "heading" else "text")
              for i, b in enumerate(raw_blocks, 1) if one_line(b.get("text", ""))]
    return ParsedDoc(title=title or "网页", doc_type="html", blocks=blocks)


# ------------------------------------------------------------------ 切块与标注
def chunk_blocks(parsed: ParsedDoc, chunk_size: int = 900) -> list[dict[str, Any]]:
    """把块合并成适合检索的段落，并自动标注公式/定义/例题。"""
    chunks: list[dict[str, Any]] = []
    buf: list[str] = []
    locs: list[str] = []
    seq = 0

    def flush() -> None:
        nonlocal seq
        if not buf:
            return
        text = "\n".join(buf).strip()
        if len(text) >= MIN_DOC_CHARS:
            kind = classify_kind(text)
            chunks.append({
                "seq": seq, "text": text, "loc": "；".join(list(dict.fromkeys(locs))[:3]) + ("…" if len(locs) > 3 else ""),
                "kind": kind, "terms": " ".join(top_terms(text, k=14)),
            })
            seq += 1
        buf.clear()
        locs.clear()

    for block in parsed.blocks:
        pieces = _split_long(block.text, chunk_size)
        for piece in pieces:
            if sum(len(x) for x in buf) + len(piece) > chunk_size and buf:
                flush()
            buf.append(piece)
            if block.loc:
                locs.append(block.loc)
        if parsed.doc_type in ("pptx", "pdf") and sum(len(x) for x in buf) > chunk_size * 0.6:
            flush()   # 课件/讲义按页聚合，检索定位更准
    flush()
    return chunks


def _split_long(text: str, size: int) -> Iterable[str]:
    text = text or ""
    if len(text) <= size:
        yield text
        return
    parts = re.split(r"(?<=[。！？；\n])", text)
    cur = ""
    for part in parts:
        if len(cur) + len(part) > size and cur:
            yield cur.strip()
            cur = ""
        cur += part
    if cur.strip():
        yield cur.strip()


def classify_kind(text: str) -> str:
    if FORMULA_HINT_RE.search(text) and len(text) < 400:
        return "formula"
    if DEF_HINT_RE.search(text):
        return "definition"
    if EXAMPLE_HINT_RE.search(text):
        return "example"
    return "text"


def top_terms(text: str, k: int = 12) -> list[str]:
    counter = Counter(t for t in tokenize(text) if len(t) > 1 and t not in STOPWORDS)
    # 专业关键词加权，保证“戴维南定理”这类术语排前面
    for word in KEY_SENTENCE_RE.findall(text or ""):
        counter[word] += 5
    return [w for w, _ in counter.most_common(k)]


def extract_formulas(text: str) -> list[str]:
    """抽取疑似公式行（含等号/希腊字母/运算符号的短行）。"""
    out = []
    for line in clean_text(text).splitlines():
        s = one_line(line)
        if 3 <= len(s) <= 160 and FORMULA_HINT_RE.search(s) and any(ch.isdigit() or ch.isalpha() for ch in s):
            out.append(s)
    return out[:20]


def extract_definitions(text: str) -> list[str]:
    out = []
    for line in clean_text(text).splitlines():
        s = one_line(line)
        if 8 <= len(s) <= 300 and DEF_HINT_RE.search(s):
            out.append(s)
    return out[:20]


def guess_knowledge_points(chunks: Sequence[dict[str, Any]], limit: int = 12) -> list[dict[str, Any]]:
    """规则法提取知识点（AI 不可用时的兜底）。"""
    points: list[dict[str, Any]] = []
    seen: set[str] = set()
    for ch in chunks:
        for formula in extract_formulas(ch.get("text", "")):
            key = normalize_for_match(formula)
            if key in seen:
                continue
            seen.add(key)
            points.append({"name": formula[:40], "type": "formula", "summary": formula,
                           "keywords": " ".join(top_terms(formula, 6)), "doc_seq": ch.get("seq")})
        for d in extract_definitions(ch.get("text", "")):
            key = normalize_for_match(d)
            if key in seen:
                continue
            seen.add(key)
            points.append({"name": d[:40], "type": "definition", "summary": d,
                           "keywords": " ".join(top_terms(d, 6)), "doc_seq": ch.get("seq")})
        if len(points) >= limit:
            break
    return points[:limit]